"""
Excel export utility for inventory items
Generates Excel file matching the trial data.xlsx format
"""

from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi.responses import StreamingResponse
import io
from datetime import datetime

from app.model.inventory import Inventory

# Reverse mapping from system units to Excel units
UNIT_REVERSE_MAPPING = {
    "Dus": "DUS",
    "Batang": "BTG",
    "Bal": "BAL",
    "Pcs": "BH",
    "Lembar": "lbr",
    "Meter": "m",
    "Pak": "pak",
    "Ons": "ons",
    "Sak": "sak",
    "Kotak": "kotak",
    "Kilogram": "kg",
    "Lusin": "ls",
}


def generate_inventory_export(inventories: List[Inventory]) -> StreamingResponse:
    """
    Generate Excel file from inventory data matching trial data.xlsx format.
    
    Format:
    - Row 1-2: Headers
    - Row 3+: Data
    - Columns: A=Kode, C=Nama, F=Satuan, G=Modal, H=Eceran, I=Qty, P=Keterangan
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    
    # Headers (Row 1 and 2) - simple version for now
    # You can customize this to match exact trial data.xlsx format
    ws['A1'] = 'Kode Barang'
    ws['C1'] = 'Nama Barang'
    ws['F1'] = 'Satuan'
    ws['G1'] = 'Harga Modal'
    ws['H1'] = 'Harga Jual Eceran'
    ws['I1'] = 'Qty'
    ws['P1'] = 'Keterangan'
    
    # Style headers
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    for col in ['A', 'C', 'F', 'G', 'H', 'I', 'P']:
        ws[f'{col}1'].font = header_font
        ws[f'{col}1'].fill = header_fill
        ws[f'{col}2'].fill = header_fill
    
    # Data rows (starting from row 3)
    row_idx = 3
    for inventory in inventories:
        # Map unit back to Excel format
        excel_unit = UNIT_REVERSE_MAPPING.get(inventory.quantity_unit, inventory.quantity_unit)
        
        ws[f'A{row_idx}'] = inventory.kode_barang
        ws[f'C{row_idx}'] = inventory.nama_barang
        ws[f'F{row_idx}'] = excel_unit
        ws[f'G{row_idx}'] = inventory.harga_modal
        ws[f'H{row_idx}'] = inventory.harga_jual_eceran
        ws[f'I{row_idx}'] = inventory.quantity
        ws[f'P{row_idx}'] = inventory.additional_note or ""
        
        row_idx += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 10
    ws.column_dimensions['P'].width = 40
    
    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Generate filename with current date
    filename = f"inventory_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    # Return as streaming response
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )
